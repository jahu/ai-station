"""Offline artifact checks. Run with the optional QA-tool environment."""

import csv
import hashlib
import json
import struct
import subprocess
import sys
import tempfile
import unittest
import zipfile
import zlib
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import operator_qa as qa
from openpyxl import load_workbook


class WorkbookTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.release = subprocess.check_output(["git", "rev-parse", "HEAD"], cwd=qa.ROOT, text=True).strip()

    def test_catalog_coverage_and_stable_unique_ids(self):
        catalog = qa.load_catalog()
        self.assertGreaterEqual(len(catalog["cases"]), 120)
        self.assertGreaterEqual(len({case["area"] for case in catalog["cases"]}), 20)
        ids = {case["id"] for case in catalog["cases"]}
        self.assertTrue(
            {"LIB-005", "LIB-006", "LIB-007", "LIB-008", "MEM-007", "PROJECT-001", "AVATAR-005"} <= ids
        )

    def test_workbook_exports_are_usable_feedback_without_prefilled_success(self):
        with tempfile.TemporaryDirectory() as folder:
            files = qa.export(
                Path(folder), release=self.release, base_url="http://127.0.0.1:8765", draft=True
            )
            self.assertEqual(len(files), 11)
            book = load_workbook(files[0], keep_links=False)
            self.assertEqual(
                book.sheetnames, ["Start", "Testy", "Zgłoszenia", "Retesty", "Podsumowanie", "Słownik"]
            )
            cases = qa.load_catalog()["cases"]
            page = book["Testy"]
            self.assertEqual([row[0].value for row in list(page.rows)[1:]], [case["id"] for case in cases])
            self.assertEqual({row[1].value for row in list(page.rows)[1:]}, {"Nie testowano"})
            self.assertEqual(page.freeze_panes, "E2")
            self.assertEqual(page.max_column, len(qa.HEADERS))
            self.assertIn(f"B2:B{len(cases) + 1}", str(page.data_validations))
            self.assertIn("WERSJA ROBOCZA", book["Start"]["B2"].value)
            instructions = dict(book["Start"].iter_rows(min_row=2, values_only=True))
            self.assertIn("Nie wybierasz trybu", instructions["Kolejność"])
            self.assertIn("suma kwota_pln = 150", instructions["Gotowe dane i oczekiwane wartości"])
            self.assertTrue(all(book["Podsumowanie"].cell(i, 2).value.endswith(")") for i in range(3, 8)))
            # Simulate a returned workbook: answers and report linkage survive save/reopen.
            page["B2"] = "Błąd"
            page["G2"] = "Zażółć: rzeczywisty błąd po wysłaniu wiadomości"
            page["R2"] = "ZGL-001"
            book["Zgłoszenia"]["C2"] = page["A2"].value
            returned = Path(folder) / "odpowiedzi.xlsx"
            book.save(returned)
            loaded = load_workbook(returned, keep_links=False)
            self.assertEqual(loaded["Testy"]["B2"].value, "Błąd")
            self.assertEqual(loaded["Testy"]["R2"].value, "ZGL-001")
            self.assertEqual(loaded["Zgłoszenia"]["C2"].value, cases[0]["id"])
            with files[1].open(encoding="utf-8-sig", newline="") as handle:
                rows = list(csv.reader(handle))
            self.assertEqual(tuple(rows[0]), qa.HEADERS)
            self.assertEqual(len(rows), len(cases) + 1)
            self.assertTrue(all(len(row) == len(qa.HEADERS) for row in rows))
            self.assertIn("\n2. ", rows[1][4])
            self.assertTrue(all(row[1] == "Nie testowano" for row in rows[1:]))
            self.assertEqual(json.loads(files[-1].read_text())["release"], self.release)

    def test_long_multiline_instructions_remain_visible_without_manual_row_resize(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            catalog = qa.load_catalog()
            case = catalog["cases"][0]
            case["steps"] = [f"Sprawdź punkt {index}." for index in range(1, 13)]
            source = root / "catalog.json"
            source.write_text(json.dumps(catalog, ensure_ascii=False))
            paths = qa.export(
                root / "export",
                release=self.release,
                base_url="http://localhost:8765",
                catalog_path=source,
                draft=True,
            )
            book = load_workbook(paths[0], keep_links=False)
            page = book["Testy"]
            self.assertGreaterEqual(page.row_dimensions[2].height, 12 * 16)
            self.assertTrue(page["E2"].alignment.wrap_text)
            self.assertIn("12. Sprawdź punkt 12.", page["E2"].value)
            self.assertEqual(page["B2"].value, "Nie testowano")
            self.assertLessEqual(max(row.height for row in page.row_dimensions.values()), 409)

    def test_synthetic_inputs_have_known_values_and_manifest_checksums(self):
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder)
            paths = qa.export(output, release=self.release, base_url="http://localhost:8765", draft=True)
            manifest = json.loads(paths[-1].read_text())
            self.assertEqual(len(manifest["fixtures"]), 5)
            for fixture in manifest["fixtures"]:
                data = (output / fixture["path"]).read_bytes()
                self.assertEqual(len(data), fixture["size_bytes"])
                self.assertEqual(hashlib.sha256(data).hexdigest(), fixture["sha256"])
            with (output / "Dane-QA/qa-sprzedaz.csv").open(encoding="utf-8-sig", newline="") as handle:
                rows = list(csv.DictReader(handle))
            self.assertEqual(len(rows), 6)
            self.assertEqual(sum(int(row["kwota_pln"]) for row in rows), 150)
            for category, total in (("A", 65), ("B", 85)):
                self.assertEqual(sum(int(r["kwota_pln"]) for r in rows if r["kategoria"] == category), total)
            book = load_workbook(output / "Dane-QA/qa-sprzedaz.xlsx", keep_links=False)
            self.assertEqual(book.sheetnames, ["Sprzedaż QA", "Kontrola QA"])
            self.assertEqual(
                dict(list(book["Kontrola QA"].values)[1:]), {"wiersze": 6, "suma": 150, "A": 65, "B": 85}
            )
            self.assertEqual(
                list(book.active.values)[1:],
                [tuple([int(r["id"]), r["produkt"], r["kategoria"], int(r["kwota_pln"])]) for r in rows],
            )
            with self.assertRaises(zipfile.BadZipFile):
                load_workbook(output / "Dane-QA/qa-uszkodzony.xlsx")
            note = (output / "Dane-QA/qa-notatka.txt").read_text()
            self.assertIn("18 listopada 2026", note)
            self.assertIn("42 uczestników", note)
            png = (output / "Dane-QA/qa-ksztalty.png").read_bytes()
            self.assertEqual(png[:8], b"\x89PNG\r\n\x1a\n")
            # Decode the complete PNG stream and check the promised source pixels.
            chunks, offset = {}, 8
            while offset < len(png):
                size = struct.unpack(">I", png[offset : offset + 4])[0]
                kind = png[offset + 4 : offset + 8]
                data = png[offset + 8 : offset + 8 + size]
                self.assertEqual(
                    struct.unpack(">I", png[offset + 8 + size : offset + 12 + size])[0],
                    zlib.crc32(kind + data),
                )
                chunks[kind] = data
                offset += 12 + size
            self.assertEqual(struct.unpack(">IIBBBBB", chunks[b"IHDR"]), (384, 256, 8, 2, 0, 0, 0))
            pixels = zlib.decompress(chunks[b"IDAT"])
            self.assertEqual(len(pixels), 256 * (1 + 384 * 3))
            for x, y, expected in (
                (96, 128, b"\xdc\x1e\x1e"),
                (272, 128, b"\x1e\x46\xdc"),
                (0, 0, b"\xff\xff\xff"),
            ):
                start = y * (1 + 384 * 3) + 1 + x * 3
                self.assertEqual(pixels[start : start + 3], expected)

    def test_existing_fixture_answers_are_preserved_before_any_export_writes(self):
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder)
            existing = output / "Dane-QA/qa-sprzedaz.csv"
            existing.parent.mkdir()
            existing.write_text("zachowaj moje odpowiedzi")
            with self.assertRaises(FileExistsError):
                qa.export(output, release=self.release, base_url="http://localhost:8765")
            self.assertEqual(existing.read_text(), "zachowaj moje odpowiedzi")
            self.assertEqual(list(output.iterdir()), [existing.parent])

    def test_refuses_overwriting_existing_operator_answers(self):
        with tempfile.TemporaryDirectory() as folder:
            paths = qa.export(Path(folder), release=self.release, base_url="http://127.0.0.1:8765")
            before = {path.name: hashlib.sha256(path.read_bytes()).hexdigest() for path in paths}
            with self.assertRaises(FileExistsError):
                qa.export(Path(folder), release=self.release, base_url="http://127.0.0.1:8765")
            self.assertEqual(
                before, {path.name: hashlib.sha256(path.read_bytes()).hexdigest() for path in paths}
            )

    def test_duplicate_case_id_is_rejected(self):
        with tempfile.TemporaryDirectory() as folder:
            catalog = qa.load_catalog()
            catalog["cases"].append(catalog["cases"][0])
            path = Path(folder) / "duplicate.json"
            path.write_text(json.dumps(catalog))
            with self.assertRaises(ValueError):
                qa.load_catalog(path)

    def test_qualified_profiles_are_readable_hashed_notes_not_operator_results(self):
        with tempfile.TemporaryDirectory() as folder:
            notes = "  LLM: Qwen 9B;\n lokalne uruchomienie.\n\nVideo: 5 sekund, 20 kroków; sprawdzony profil QA.  "
            files = qa.export(
                Path(folder),
                release=self.release,
                base_url="http://localhost:8765",
                draft=True,
                qualification_notes=notes,
            )
            book = load_workbook(files[0], keep_links=False)
            rows = {row[0].value: row for row in book["Start"].iter_rows(min_row=2)}
            expected = [
                "LLM: Qwen 9B; lokalne uruchomienie.",
                "Video: 5 sekund, 20 kroków; sprawdzony profil QA.",
            ]
            for index, text in enumerate(expected, 1):
                row = rows[f"Sprawdzony profil {index}"]
                self.assertEqual(row[1].value, text)
                self.assertTrue(row[1].alignment.wrap_text)
                self.assertGreaterEqual(book["Start"].row_dimensions[row[0].row].height, 32)
            self.assertEqual({row[1] for row in list(book["Testy"].values)[1:]}, {"Nie testowano"})
            guide = next(p for p in files if p.name.endswith("-Instrukcja.txt")).read_text()
            self.assertTrue(all(text in guide for text in expected))
            manifest = json.loads(files[-1].read_text())
            self.assertEqual(manifest["qualification_paragraphs"], 2)
            self.assertEqual(
                manifest["qualification_notes_sha256"],
                hashlib.sha256("\n\n".join(expected).encode()).hexdigest(),
            )

    def test_unprovided_or_formula_like_profiles_do_not_imply_verified_execution(self):
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder)
            files = qa.export(
                output / "empty", release=self.release, base_url="http://localhost:8765", draft=True
            )
            manifest = json.loads(files[-1].read_text())
            self.assertIsNone(manifest["qualification_notes_sha256"])
            self.assertEqual(manifest["qualification_paragraphs"], 0)
            info = dict(load_workbook(files[0])["Start"].iter_rows(min_row=2, values_only=True))
            self.assertIn("nie został dołączony", info["Zakres sprawdzonych profili"])
            files = qa.export(
                output / "literal",
                release=self.release,
                base_url="http://localhost:8765",
                draft=True,
                qualification_notes='=HYPERLINK("https://example.invalid","literal")',
            )
            cell = next(
                row[1]
                for row in load_workbook(files[0])["Start"].iter_rows(min_row=2)
                if row[0].value == "Sprawdzony profil 1"
            )
            self.assertEqual(cell.data_type, "s")
            self.assertTrue(cell.value.startswith("="))

    def test_unreadably_large_profile_notes_are_rejected_before_export_writes(self):
        for notes in ("x" * 32769, "x" * 1201, "\n\n".join(["profile"] * 25)):
            with self.subTest(notes_length=len(notes)), tempfile.TemporaryDirectory() as folder:
                output = Path(folder) / "export"
                with self.assertRaises(ValueError):
                    qa.export(
                        output,
                        release=self.release,
                        base_url="http://localhost:8765",
                        qualification_notes=notes,
                    )
                self.assertFalse(output.exists())

    def test_catalog_text_never_becomes_an_excel_or_csv_formula(self):
        with tempfile.TemporaryDirectory() as folder:
            catalog = qa.load_catalog()
            catalog["cases"][0]["title"] = '=HYPERLINK("https://example.invalid","untrusted")'
            path = Path(folder) / "input.json"
            path.write_text(json.dumps(catalog))
            files = qa.export(
                Path(folder), release=self.release, base_url="http://localhost:8765", catalog_path=path
            )
            cell = load_workbook(files[0])["Testy"]["D2"]
            self.assertEqual(cell.data_type, "s")
            with files[1].open(encoding="utf-8-sig", newline="") as handle:
                rows = list(csv.reader(handle))
            self.assertTrue(rows[1][3].startswith("'="))

    def test_release_and_ui_address_must_not_hide_moving_refs_or_credentials(self):
        with tempfile.TemporaryDirectory() as folder:
            for release, url in [
                ("develop", "http://localhost:8765"),
                (self.release, "https://user:secret@example.invalid"),
            ]:
                with self.assertRaises(ValueError):
                    qa.export(Path(folder), release=release, base_url=url)


if __name__ == "__main__":
    unittest.main()
