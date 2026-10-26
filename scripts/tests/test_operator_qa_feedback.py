"""Offline feedback acceptance; synthetic operator answers never touch the app."""

import csv
import hashlib
import json
import os
import re
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import operator_qa as qa
import operator_qa_feedback as feedback
from openpyxl import load_workbook


class FeedbackTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.release = subprocess.check_output(["git", "rev-parse", "HEAD"], cwd=qa.ROOT, text=True).strip()

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.paths = qa.export(
            self.root / "package", release=self.release, base_url="http://127.0.0.1:8765", draft=True
        )
        self.manifest = self.paths[-1]
        self.book = load_workbook(self.paths[0])
        self.returned = self.root / "odpowiedzi-zażółć.xlsx"
        self.case_id = self.book["Testy"]["A2"].value

    def ingest(self, **changes):
        self.book.save(self.returned)
        return feedback.ingest(
            **{
                "manifest_path": self.manifest,
                "workbook": self.returned,
                "output": self.root / "triage",
                **changes,
            }
        )

    def issue(self, number, *, kind="Błąd", title="Wysłanie wiadomości", actual="Błąd po Send"):
        values = [
            f"ZGL-{number:03}",
            kind,
            self.case_id,
            title,
            "1. Wpisz\n2. Wyślij",
            "Odpowiedź",
            actual,
            "Zawsze",
            "Wysoki",
            f"run:{number}",
            "telefon / model",
            "Nowe",
            "",
        ]
        for col, value in enumerate(values, 1):
            self.book["Zgłoszenia"].cell(number + 1, col, value)

    def test_blank_package_stays_untested_and_preserves_originals_without_overwrite(self):
        with patch("socket.create_connection", side_effect=AssertionError("No network")):
            result = self.ingest()
        self.assertEqual(result["counts"], {"Nie testowano": len(qa.load_catalog()["cases"])})
        self.assertEqual(result["triage"], [])
        self.assertEqual(result["retests"], [])
        self.assertTrue(result["source_package_draft"])
        before = self.returned.read_bytes()
        self.assertEqual((self.root / "triage/originals/returned.xlsx").read_bytes(), before)
        self.assertEqual(result["originals"]["returned.xlsx"], hashlib.sha256(before).hexdigest())
        self.assertEqual(os.stat(self.root / "triage").st_mode & 0o777, 0o700)
        self.assertEqual(os.stat(self.root / "triage/triage.json").st_mode & 0o777, 0o600)
        with self.assertRaises(FileExistsError):
            feedback.ingest(manifest_path=self.manifest, workbook=self.returned, output=self.root / "triage")
        self.assertEqual(self.returned.read_bytes(), before)

    def test_reports_group_only_same_symptoms_keep_evidence_and_retests_never_erase_failure(self):
        self.book["Testy"]["B2"] = "Błąd"
        self.book["Testy"]["G2"] = "Zażółć gęślą jaźń — brak odpowiedzi"
        self.book["Testy"]["R2"] = "ZGL-001; ZGL-002"
        self.book["Testy"]["B3"] = "Zablokowany"
        self.issue(1)
        self.issue(2)
        self.issue(3, kind="Usprawnienie", title="Łatwiejszy wybór modelu")
        self.issue(4, actual="Odpowiedź jest ucięta")
        for col, value in enumerate(
            ["ZGL-001", self.case_id, "fix-123", "2026-10-26", "Send", "OK", "", "run:5"], 1
        ):
            self.book["Retesty"].cell(2, col, value)
        result = self.ingest()
        self.assertEqual(result["results"][0]["Wynik"], "Błąd")
        self.assertEqual(result["retests"][0]["Wynik retestu"], "OK")
        self.assertEqual(len(result["reports"]), 4)
        duplicates = next(row for row in result["triage"] if len(row["members"]) == 2)
        self.assertEqual([row["id"] for row in duplicates["members"]], ["ZGL-001", "ZGL-002"])
        self.assertEqual(
            [row["source"]["Dowód / ID wykonania"] for row in duplicates["members"]], ["run:1", "run:2"]
        )
        self.assertEqual(
            [row["kind"] for row in result["triage"]], ["Błąd", "Błąd", "Blokada", "Usprawnienie"]
        )
        self.assertTrue(result["warnings"])
        self.assertIn("Zażółć", (self.root / "triage/triage.json").read_text())

    def test_improvement_link_does_not_hide_failed_case(self):
        self.book["Testy"]["B2"] = "Błąd"
        self.book["Testy"]["R2"] = "ZGL-001"
        self.issue(1, kind="Usprawnienie")
        result = self.ingest()
        self.assertEqual([item["kind"] for item in result["triage"]], ["Błąd", "Usprawnienie"])

    def test_incomplete_report_does_not_hide_richer_case_result(self):
        self.book["Testy"]["B2"] = "Błąd"
        self.book["Testy"]["G2"] = "Szczegółowy wynik bez odpowiedzi"
        self.book["Testy"]["R2"] = "ZGL-001"
        self.issue(1, actual="")
        result = self.ingest()
        self.assertEqual(len(result["triage"]), 2)
        self.assertFalse(result["triage"][0]["needs_clarification"])
        self.assertTrue(result["triage"][1]["needs_clarification"])
        self.assertTrue(
            any(item["actual"] == "Szczegółowy wynik bez odpowiedzi" for item in result["triage"])
        )

    def test_csv_answers_are_literal_inert_text_and_need_all_three_tables(self):
        sources = {}
        for name, headers in feedback.TABLES.items():
            page = self.book[name]
            if name == "Testy":
                page["B2"] = "Błąd"
                page["G2"] = '=HYPERLINK("https://invalid.example/","tekst")'
                page["G2"].data_type = "s"
            target = self.root / (name + ".csv")
            with target.open("w", encoding="utf-8-sig", newline="") as handle:
                csv.writer(handle).writerows(page.values)
            sources[name] = target
        result = feedback.ingest(manifest_path=self.manifest, csv_paths=sources, output=self.root / "csv")
        self.assertTrue(result["results"][0]["Rzeczywisty rezultat"].startswith("=HYPERLINK"))
        self.assertEqual(result["identity_binding"], "explicit CSV manifest")
        self.assertEqual(result["counts"]["Błąd"], 1)
        with self.assertRaisesRegex(ValueError, "three"):
            feedback.ingest(
                manifest_path=self.manifest,
                csv_paths={"Testy": sources["Testy"]},
                output=self.root / "missing",
            )

    def test_invalid_headers_ids_status_and_changed_definition_fail_before_output(self):
        for cell, value in [
            ("A1", "wrong"),
            ("A2", "UNKNOWN-001"),
            ("A3", self.case_id),
            ("B2", "Maybe"),
            ("D2", "Changed definition"),
        ]:
            with self.subTest(cell=cell):
                self.book = load_workbook(self.paths[0])
                self.book["Testy"][cell] = value
                with self.assertRaises(ValueError):
                    self.ingest()
                self.assertFalse((self.root / "triage").exists())
        self.book = load_workbook(self.paths[0])
        self.book["Testy"].delete_rows(2)
        with self.assertRaisesRegex(ValueError, "Missing case"):
            self.ingest()

    def test_manifest_catalog_and_workbook_identity_cannot_be_substituted(self):
        for key, value in [
            ("release", "a" * 40),
            ("catalog_version", 999),
            ("catalog_sha256", "b" * 64),
            ("case_ids", []),
            ("schema_version", True),
        ]:
            with self.subTest(key=key):
                manifest = json.loads(self.manifest.read_text())
                manifest[key] = value
                target = self.root / "bad-manifest.json"
                target.write_text(json.dumps(manifest))
                with self.assertRaises(ValueError):
                    self.ingest(manifest_path=target)
        self.book["Start"]["B3"] = "b" * 40
        with self.assertRaisesRegex(ValueError, "identity mismatch"):
            self.ingest()

    def test_formulas_in_answer_or_identity_sheets_are_rejected_but_summary_is_not_executed(self):
        for name, cell in [("Testy", "G2"), ("Zgłoszenia", "D2"), ("Retesty", "G2"), ("Start", "B2")]:
            with self.subTest(sheet=name):
                self.book = load_workbook(self.paths[0])
                self.book[name][cell] = '=HYPERLINK("https://invalid.example/","never run")'
                with self.assertRaisesRegex(ValueError, "Formula/error"):
                    self.ingest()
        self.book = load_workbook(self.paths[0])
        self.book["Podsumowanie"]["B3"] = '=HYPERLINK("https://invalid.example/","ignored")'
        result = self.ingest()
        self.assertEqual(result["counts"]["Nie testowano"], len(qa.load_catalog()["cases"]))

    def altered_archive(self, name, content):
        target = self.root / "archive.xlsx"
        with ZipFile(self.paths[0]) as original, ZipFile(target, "w") as altered:
            for item in original.infolist():
                if item.filename != name:
                    altered.writestr(item.filename, original.read(item))
            altered.writestr(name, content)
        return target

    def test_macros_external_workbooks_entity_xml_and_archive_size_are_rejected(self):
        for name, content in [
            ("xl/vbaProject.bin", b"untrusted"),
            ("xl/externalLinks/externalLink1.xml", b"untrusted"),
            ("extra.xml", b'<!DOCTYPE x [<!ENTITY y "x">]><x>&y;</x>'),
            ("extra.xml", "<!DOCTYPE x><x/>".encode("utf-16")),
        ]:
            with self.subTest(name=name):
                with self.assertRaises(ValueError):
                    feedback.ingest(
                        manifest_path=self.manifest,
                        workbook=self.altered_archive(name, content),
                        output=self.root / "triage",
                    )
        with patch.object(feedback, "MAX_UNPACKED", 1):
            with self.assertRaisesRegex(ValueError, "archive exceeds"):
                self.ingest()
        with patch.object(feedback, "MAX_BYTES", 1):
            with self.assertRaisesRegex(ValueError, "exceeds"):
                self.ingest()

    def test_forged_dimensions_do_not_hide_rows(self):
        with ZipFile(self.paths[0]) as archive:
            raw = archive.read("xl/worksheets/sheet2.xml")
        forged = re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1:A1"', raw)
        source = self.altered_archive("xl/worksheets/sheet2.xml", forged)
        result = feedback.ingest(manifest_path=self.manifest, workbook=source, output=self.root / "triage")
        self.assertEqual(len(result["results"]), len(qa.load_catalog()["cases"]))

    def test_duplicate_zip_members_and_macro_enabled_content_type_reject(self):
        target = self.root / "duplicate.xlsx"
        with ZipFile(self.paths[0]) as source, ZipFile(target, "w") as duplicate:
            for entry in source.infolist():
                duplicate.writestr(entry.filename, source.read(entry))
            with self.assertWarns(UserWarning):
                duplicate.writestr("xl/worksheets/sheet2.xml", source.read("xl/worksheets/sheet2.xml"))
        with self.assertRaisesRegex(ValueError, "duplicate workbook"):
            feedback.ingest(manifest_path=self.manifest, workbook=target, output=self.root / "duplicate")
        with ZipFile(self.paths[0]) as source:
            content = source.read("[Content_Types].xml").replace(
                b"spreadsheetml.sheet.main+xml", b"ms-excel.sheet.macroEnabled.main+xml"
            )
        target = self.altered_archive("[Content_Types].xml", content)
        with self.assertRaisesRegex(ValueError, "Macro-enabled"):
            feedback.ingest(manifest_path=self.manifest, workbook=target, output=self.root / "macro")

    def test_unknown_issue_links_warn_but_unknown_cases_or_duplicate_reports_reject(self):
        self.book["Testy"]["B2"] = "Błąd"
        self.book["Testy"]["R2"] = "EARLIER-123"
        self.book["Retesty"]["A2"] = "EARLIER-123"
        self.book["Retesty"]["B2"] = self.case_id
        self.book["Retesty"]["F2"] = "OK"
        result = self.ingest()
        self.assertEqual(len(result["triage"]), 1)
        self.assertTrue(any("EARLIER-123" in warning for warning in result["warnings"]))
        self.book = load_workbook(self.paths[0])
        self.issue(1)
        self.book["Zgłoszenia"]["C2"] = "UNKNOWN-001"
        with self.assertRaisesRegex(ValueError, "Unknown case"):
            self.ingest(output=self.root / "unknown")
        self.book["Zgłoszenia"]["C2"] = self.case_id
        self.issue(2)
        self.book["Zgłoszenia"]["A3"] = "ZGL-001"
        with self.assertRaisesRegex(ValueError, "duplicate report"):
            self.ingest(output=self.root / "duplicate")

    def test_unclassified_reports_and_html_remain_data_not_instructions(self):
        self.issue(1, kind="", title='<script>alert("x")</script> [link](https://invalid.example)')
        self.book["Zgłoszenia"]["C2"] = ""  # Observations outside the catalog are supported.
        result = self.ingest()
        self.assertEqual(result["triage"][0]["kind"], "Do wyjaśnienia")
        rendered = (self.root / "triage/TRIAGE.md").read_text()
        self.assertNotIn("<script>", rendered)
        self.assertIn("&lt;script&gt;", rendered)
        self.assertIn(r"\[link\]", rendered)
        self.assertTrue(result["warnings"])

    def test_cli_writes_reproducible_bundle_from_a_returned_workbook(self):
        self.book.save(self.returned)
        result = subprocess.run(
            [
                sys.executable,
                str(qa.ROOT / "scripts/operator_qa_feedback.py"),
                "--manifest",
                str(self.manifest),
                "--workbook",
                str(self.returned),
                "--output",
                str(self.root / "cli"),
            ],
            capture_output=True,
            text=True,
            timeout=30,
        )
        self.assertEqual(result.returncode, 0, result.stderr)
        saved = json.loads((self.root / "cli/triage.json").read_text())
        self.assertEqual(saved["release"], self.release)
        self.assertEqual(saved["triage"], [])


if __name__ == "__main__":
    unittest.main()
