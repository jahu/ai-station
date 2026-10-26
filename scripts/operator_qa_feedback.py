"""Read returned QA answers as inert data and produce an offline triage bundle."""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import re
import subprocess
from collections import Counter
from datetime import UTC, date, datetime
from html import escape
from pathlib import Path
from zipfile import ZipFile

import operator_qa as qa

MAX_BYTES = 32 * 1024 * 1024
MAX_UNPACKED = 64 * 1024 * 1024
MAX_ROWS = 10000
TABLES = {"Testy": qa.HEADERS, "Zgłoszenia": qa.ISSUE_HEADERS, "Retesty": qa.RETEST_HEADERS}


def read_bytes(path):
    with Path(path).open("rb") as handle:
        value = handle.read(MAX_BYTES + 1)
    if len(value) > MAX_BYTES:
        raise ValueError(f"Input exceeds {MAX_BYTES} bytes: {path}")
    return value


def text(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return "" if value is None else str(value)


def references(value):
    return list(dict.fromkeys(re.split(r"[;,\s]+", value.strip()))) if value.strip() else []


def table(rows, headers, name):
    rows = iter(rows)
    if tuple(next(rows, ())) != tuple(headers):
        raise ValueError(f"Invalid {name} headers")
    result = []
    for number, row in enumerate(rows, 2):
        if number > MAX_ROWS:
            raise ValueError(f"Too many {name} rows")
        if not any(row):
            continue
        if len(row) != len(headers):
            raise ValueError(f"Invalid {name} column count at row {number}")
        if any(len(value) > 32767 for value in row):
            raise ValueError(f"Oversized {name} cell at row {number}")
        result.append(dict(zip(headers, row, strict=True)))
    return result


def workbook_tables(raw):
    from openpyxl import load_workbook

    with ZipFile(io.BytesIO(raw)) as archive:
        entries = archive.infolist()
        if len({entry.filename for entry in entries}) != len(entries):
            raise ValueError("Ambiguous duplicate workbook archive entries")
        if len(entries) > 2000 or sum(entry.file_size for entry in entries) > MAX_UNPACKED:
            raise ValueError("Workbook archive exceeds import bounds")
        for entry in entries:
            name = entry.filename.lower()
            if name.endswith("vbaproject.bin") or name.startswith("xl/externallinks/"):
                raise ValueError("Macros and external workbook links are not accepted")
            if name.endswith((".xml", ".rels")):
                content = archive.read(entry).upper().replace(b"\0", b"")
                if name == "[content_types].xml" and b"MACROENABLED" in content:
                    raise ValueError("Macro-enabled workbooks are not accepted")
                if b"<!DOCTYPE" in content or b"<!ENTITY" in content:
                    raise ValueError("XML declarations with entities are not accepted")
    book = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
    try:

        def values(name):
            if name not in book.sheetnames:
                raise ValueError(f"Missing sheet: {name}")
            page = book[name]
            page.reset_dimensions()  # Do not trust a forged, truncated worksheet dimension.
            for number, row in enumerate(page.iter_rows(), 1):
                if number > MAX_ROWS or len(row) > 64:
                    raise ValueError(f"Oversized sheet: {name}")
                if any(cell.data_type in {"f", "e"} for cell in row):
                    raise ValueError(f"Formula/error in imported sheet {name}, row {number}")
                yield [text(cell.value) for cell in row]

        start = table(values("Start"), ("Informacja", "Instrukcja"), "Start")
        if len({row["Informacja"] for row in start}) != len(start):
            raise ValueError("Duplicate Start identity labels")
        identity = {row["Informacja"]: row["Instrukcja"] for row in start}
        return {name: table(values(name), headers, name) for name, headers in TABLES.items()}, identity
    finally:
        book.close()


def checked_manifest(raw, catalog):
    manifest = json.loads(raw)
    if not isinstance(manifest, dict):
        raise ValueError("Manifest must be an object")
    digest = hashlib.sha256(json.dumps(catalog, sort_keys=True, ensure_ascii=False).encode()).hexdigest()
    if (
        type(manifest.get("schema_version")) is not int
        or manifest["schema_version"] != 1
        or manifest.get("catalog_version") != catalog["version"]
        or manifest.get("catalog_sha256") != digest
        or manifest.get("case_ids") != [case["id"] for case in catalog["cases"]]
        or not re.fullmatch(r"[0-9a-f]{40}", text(manifest.get("release")))
        or not isinstance(manifest.get("draft"), bool)
    ):
        raise ValueError("Manifest does not match the reviewed catalog/release schema")
    probe = subprocess.run(
        ["git", "cat-file", "-t", manifest["release"]], cwd=qa.ROOT, capture_output=True, text=True
    )
    if probe.returncode or probe.stdout.strip() != "commit":
        raise ValueError("Manifest release is not an available immutable Git commit")
    return manifest


def compile_feedback(tables, catalog, *, csv_input):
    expected = {row[0]: dict(zip(qa.HEADERS, map(text, row), strict=True)) for row in qa.rows(catalog)}
    cases, warnings = {}, []
    fixed = [qa.HEADERS[index] for index in (0, 2, 3, 4, 5, 10, 11, 12, 13, 18)]
    for row in tables["Testy"]:
        key = row["ID"]
        if key not in expected or key in cases:
            raise ValueError(f"Unknown or duplicate case ID: {key}")
        for name in fixed:
            value = expected[key][name]
            if csv_input and value.lstrip().startswith(("=", "+", "-", "@")):
                value = "'" + value
            if row[name] != value:
                raise ValueError(f"Changed case definition: {key}, {name}")
        if row["Wynik"] not in qa.STATUSES:
            raise ValueError(f"Invalid result for {key}: {row['Wynik']}")
        if row["Wynik"] in {"Zablokowany", "Nie dotyczy"} and not row["Uwagi / powód blokady"].strip():
            warnings.append(f"{key}: brak uzasadnienia wyniku {row['Wynik']}")
        if row["Wynik"] == "Błąd" and not row["Rzeczywisty rezultat"].strip():
            warnings.append(f"{key}: brak opisu rzeczywistego wyniku; zachowano Błąd")
        cases[key] = row
    if cases.keys() != expected.keys():
        raise ValueError("Missing case rows; retain every stable ID from the manifest")

    issues, linked = {}, set()
    for row in tables["Zgłoszenia"]:
        if not any(row[name].strip() for name in qa.ISSUE_HEADERS[1:]):
            continue  # Exported blank slots already have a report ID.
        key = row["ID zgłoszenia"]
        if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9._-]{0,79}", key) or key in issues:
            raise ValueError(f"Invalid or duplicate report ID: {key}")
        case_ids = references(row["ID testów"])
        if set(case_ids) - cases.keys():
            raise ValueError(f"Unknown case reference in report {key}")
        if row["Typ"] not in {"", "Błąd", "Usprawnienie"}:
            raise ValueError(f"Invalid report type: {key}")
        if row["Waga"] not in {"", "Krytyczny", "Wysoki", "Średni", "Niski"}:
            raise ValueError(f"Invalid severity: {key}")
        issues[key] = {"id": key, "case_ids": case_ids, "source": row}
    for key, row in cases.items():
        for issue_id in references(row["ID zgłoszeń"]):
            if issue_id not in issues:
                warnings.append(
                    f"{key}: nie dołączono zgłoszenia {issue_id}; zachowano oryginalne powiązanie"
                )
            elif key not in issues[issue_id]["case_ids"]:
                issues[issue_id]["case_ids"].append(key)
    findings = []
    for issue in issues.values():
        row = issue["source"]
        if row["Typ"] == "Błąd" and row["Rzeczywiste"].strip():
            linked.update(issue["case_ids"])
        if (
            not row["Typ"]
            or not row["Tytuł"].strip()
            or (row["Typ"] == "Błąd" and not row["Rzeczywiste"].strip())
        ):
            warnings.append(f"{issue['id']}: niepełne zgłoszenie zachowano do wyjaśnienia")
        findings.append(
            {
                "id": issue["id"],
                "kind": row["Typ"] or "Do wyjaśnienia",
                "title": row["Tytuł"] or issue["id"],
                "case_ids": sorted(issue["case_ids"]),
                "severity": row["Waga"],
                "steps": row["Kroki odtworzenia"],
                "expected": row["Oczekiwane"],
                "actual": row["Rzeczywiste"],
                "source": row,
            }
        )
    for key, row in cases.items():
        if row["Wynik"] == "Zablokowany" or (row["Wynik"] == "Błąd" and key not in linked):
            findings.append(
                {
                    "id": "case:" + key,
                    "kind": "Błąd" if row["Wynik"] == "Błąd" else "Blokada",
                    "title": row["Cel testu"],
                    "case_ids": [key],
                    "severity": "",
                    "steps": row["Kroki"],
                    "expected": row["Oczekiwany wynik"],
                    "actual": row["Rzeczywisty rezultat"] or row["Uwagi / powód blokady"],
                    "source": row,
                }
            )
    retests = []
    for row in tables["Retesty"]:
        if not any(value.strip() for value in row.values()):
            continue
        if row["ID testu"] and row["ID testu"] not in cases:
            raise ValueError("Unknown retest case ID")
        if row["Wynik retestu"] not in qa.STATUSES or not row["ID zgłoszenia"]:
            raise ValueError("Retest needs a report ID and valid result")
        if row["ID zgłoszenia"] not in issues:
            warnings.append(
                f"Retest dotyczy wcześniejszego lub niedołączonego zgłoszenia: {row['ID zgłoszenia']}"
            )
        retests.append(row)
    groups = {}
    for finding in findings:
        # No fuzzy merge: different observations or case scopes remain separate.
        key = tuple(
            " ".join(text(finding[field]).split()).casefold()
            for field in ("kind", "title", "case_ids", "severity", "steps", "expected", "actual")
        )
        groups.setdefault(key, []).append(finding)
    ranked = []
    for group in groups.values():
        lead = group[0]
        priorities = [cases[key]["Priorytet"] for key in lead["case_ids"]]
        ranked.append(
            {
                **{key: value for key, value in lead.items() if key != "source"},
                "priority": min(priorities, default=""),
                "members": group,
            }
        )
    ranked.sort(
        key=lambda row: (
            {"Błąd": 0, "Blokada": 1, "Usprawnienie": 2}.get(row["kind"], 3),
            {"Krytyczny": 0, "Wysoki": 1, "Średni": 2, "Niski": 3}.get(row["severity"], 4),
            row["priority"] or "P9",
            row["id"],
        )
    )
    return {
        "results": list(cases.values()),
        "reports": list(issues.values()),
        "retests": retests,
        "counts": dict(Counter(row["Wynik"] for row in cases.values())),
        "triage": ranked,
        "warnings": warnings,
    }


def markdown(result):
    def safe(value):
        return re.sub(r"([\\`*_{}\[\]#|])", r"\\\1", escape(text(value)))

    lines = [
        "# Wyniki QA — lista do przeglądu",
        "",
        f"Wersja aplikacji: {result['release']}",
        "",
        "Dane operatora wymagają przeglądu. Import nie uruchamia napraw ani nie zmienia wyników.",
        "",
        "Wyniki: " + "; ".join(f"{key}: {value}" for key, value in result["counts"].items()),
        "",
    ]
    for index, finding in enumerate(result["triage"], 1):
        lines += [
            f"## {index}. {safe(finding['kind'])}: {safe(finding['title'])}",
            "",
            "ID źródłowe: " + ", ".join(safe(row["id"]) for row in finding["members"]),
            "Testy: " + (", ".join(finding["case_ids"]) or "obserwacja poza katalogiem"),
            f"Waga: {safe(finding['severity']) or 'nie podano'}; priorytet testu: {finding['priority'] or 'brak'}",
            "",
        ]
        for label, field in (("Kroki", "steps"), ("Oczekiwane", "expected"), ("Rzeczywiste", "actual")):
            lines += [label + ":", *("> " + safe(line) for line in finding[field].splitlines()), ""]
        for member in finding["members"]:
            source = member["source"]
            evidence = [
                source.get(name, "")
                for name in (
                    "Dowód / ID wykonania",
                    "Dowód / link",
                    "ID Run / rozmowy / agenta / pliku",
                    "Wersja / urządzenie / model",
                    "Model / profil",
                    "Urządzenie / przeglądarka",
                )
            ]
            lines += [
                "Dowód / kontekst "
                + safe(member["id"])
                + ": "
                + "; ".join(safe(value) for value in evidence if value),
                "",
            ]
    lines += ["## Uwagi do wyjaśnienia", "", *("- " + safe(line) for line in result["warnings"])]
    lines += [
        "",
        f"Osobno zachowano retesty: {len(result['retests'])}. Pierwotne wyniki pozostają bez zmian.",
        "",
    ]
    return "\n".join(lines)


def ingest(*, manifest_path, output, catalog_path=qa.CATALOG, workbook=None, csv_paths=None):
    if (workbook is None) == (csv_paths is None):
        raise ValueError("Choose one workbook or all three CSV files")
    originals = {"manifest.json": read_bytes(manifest_path), "catalog.json": read_bytes(catalog_path)}
    catalog = qa.validate_catalog(json.loads(originals["catalog.json"]))
    manifest = checked_manifest(originals["manifest.json"], catalog)
    if workbook is not None:
        originals["returned.xlsx"] = read_bytes(workbook)
        tables, identity = workbook_tables(originals["returned.xlsx"])
        for name, value in (
            ("Wersja aplikacji", manifest["release"]),
            ("Wersja katalogu", str(manifest["catalog_version"])),
            ("Identyfikator katalogu", manifest["catalog_sha256"]),
        ):
            if identity.get(name) != value:
                raise ValueError(f"Workbook/manifest identity mismatch: {name}")
    else:
        if set(csv_paths) != set(TABLES):
            raise ValueError("Supply all three CSV files: Testy, Zgłoszenia and Retesty")
        tables = {}
        for name, path in csv_paths.items():
            originals[name + ".csv"] = raw = read_bytes(path)
            tables[name] = table(
                csv.reader(io.StringIO(raw.decode("utf-8-sig")), strict=True), TABLES[name], name
            )
    result = {
        "schema_version": 1,
        "release": manifest["release"],
        "catalog_version": manifest["catalog_version"],
        "catalog_sha256": manifest["catalog_sha256"],
        "source_package_draft": manifest["draft"],
        "generated_at": datetime.now(UTC).isoformat(),
        "identity_binding": "workbook Start and explicit manifest" if workbook else "explicit CSV manifest",
        **compile_feedback(tables, catalog, csv_input=workbook is None),
    }
    # Validate everything before creating a new bundle. Input bytes are never rewritten.
    output = Path(output)
    output.mkdir(mode=0o700, parents=True, exist_ok=False)
    (output / "originals").mkdir(mode=0o700)
    for name, raw in originals.items():
        (output / "originals" / name).write_bytes(raw)
        (output / "originals" / name).chmod(0o600)
    result["originals"] = {name: hashlib.sha256(raw).hexdigest() for name, raw in originals.items()}
    (output / "triage.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    (output / "TRIAGE.md").write_text(markdown(result), encoding="utf-8")
    for path in (output / "triage.json", output / "TRIAGE.md"):
        path.chmod(0o600)
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--manifest", required=True, type=Path)
    parser.add_argument("--catalog", type=Path, default=qa.CATALOG)
    source = parser.add_mutually_exclusive_group(required=True)
    source.add_argument("--workbook", type=Path)
    source.add_argument("--tests-csv", type=Path)
    parser.add_argument("--issues-csv", type=Path)
    parser.add_argument("--retests-csv", type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    if args.workbook and (args.issues_csv or args.retests_csv):
        parser.error("Workbook cannot be combined with CSV inputs")
    if args.tests_csv and not (args.issues_csv and args.retests_csv):
        parser.error("CSV import needs all three files and their original manifest")
    ingest(
        manifest_path=args.manifest,
        catalog_path=args.catalog,
        output=args.output,
        workbook=args.workbook,
        csv_paths={"Testy": args.tests_csv, "Zgłoszenia": args.issues_csv, "Retesty": args.retests_csv}
        if args.tests_csv
        else None,
    )
    print(f"Saved inert feedback and triage: {args.output}")


if __name__ == "__main__":
    main()
