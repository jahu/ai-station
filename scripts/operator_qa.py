"""Export the Polish operator case catalog; never execute workloads or test results."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import subprocess
from datetime import UTC, datetime
from pathlib import Path
from textwrap import wrap as wrap_text
from urllib.parse import urlsplit

from operator_qa_fixtures import FIXTURE_INSTRUCTIONS, fixture_bytes

ROOT = Path(__file__).resolve().parent.parent
CATALOG = ROOT / "docs/qa/cases.json"
STATUSES = ("Nie testowano", "OK", "Błąd", "Zablokowany", "Nie dotyczy")
HEADERS = (
    "ID",
    "Wynik",
    "Obszar",
    "Cel testu",
    "Kroki",
    "Oczekiwany wynik",
    "Rzeczywisty rezultat",
    "Uwagi / powód blokady",
    "Dowód / link",
    "ID Run / rozmowy / agenta / pliku",
    "Priorytet",
    "Warunki i dane",
    "Zasoby / koszty",
    "Minuty",
    "Model / profil",
    "Urządzenie / przeglądarka",
    "Data wykonania",
    "ID zgłoszeń",
    "Wymagania",
)
ISSUE_HEADERS = (
    "ID zgłoszenia",
    "Typ",
    "ID testów",
    "Tytuł",
    "Kroki odtworzenia",
    "Oczekiwane",
    "Rzeczywiste",
    "Częstotliwość",
    "Waga",
    "Dowód / ID wykonania",
    "Wersja / urządzenie / model",
    "Stan triage",
    "Poprawka / retest",
)
RETEST_HEADERS = (
    "ID zgłoszenia",
    "ID testu",
    "Wersja poprawki",
    "Data",
    "Kroki / zakres",
    "Wynik retestu",
    "Uwagi",
    "Dowód",
)


def load_catalog(path: Path = CATALOG) -> dict:
    return validate_catalog(json.loads(path.read_text(encoding="utf-8")))


def validate_catalog(value: dict) -> dict:
    cases = value["cases"]
    ids = [item["id"] for item in cases]
    if not cases or len(ids) != len(set(ids)):
        raise ValueError("Case IDs must be nonempty and unique")
    for case in cases:
        if not re.fullmatch(r"[A-Z]+-\d{3}", case["id"]):
            raise ValueError("Invalid stable case ID")
        for field in ("area", "title", "priority", "resources", "requirements"):
            if not case.get(field):
                raise ValueError(f"Missing {field} in {case['id']}")
        for field in ("prerequisites", "steps", "expected", "requirements"):
            if (
                not isinstance(case.get(field), list)
                or not case[field]
                or not all(isinstance(item, str) and item.strip() for item in case[field])
            ):
                raise ValueError(f"Missing {field} in {case['id']}")
        if case["priority"] not in {"P0", "P1", "P2"} or not 1 <= case["minutes"] <= 180:
            raise ValueError(f"Invalid workload/priority in {case['id']}")
    return value


def rows(catalog):
    for case in catalog["cases"]:
        yield [
            case["id"],
            STATUSES[0],
            case["area"],
            case["title"],
            "\n".join(f"{i}. {step}" for i, step in enumerate(case["steps"], 1)),
            "\n".join(case["expected"]),
            *([""] * 4),
            case["priority"],
            "\n".join(case["prerequisites"]),
            case["resources"],
            case["minutes"],
            *([""] * 4),
            "; ".join(case["requirements"]),
        ]


def export(
    output: Path,
    *,
    release: str,
    base_url: str,
    draft=False,
    catalog_path=CATALOG,
    qualification_notes: str = "",
):
    """Only explicit immutable release identities can label a handoff workbook."""
    from openpyxl import Workbook
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    if not re.fullmatch(r"[0-9a-f]{40}", release):
        raise ValueError("Release must be a full immutable Git commit SHA")
    if subprocess.check_output(["git", "cat-file", "-t", release], cwd=ROOT, text=True).strip() != "commit":
        raise ValueError("Release must name a commit in this repository")
    target = urlsplit(base_url)
    if (
        target.scheme not in {"http", "https"}
        or not target.hostname
        or target.username
        or target.password
        or target.query
        or target.fragment
    ):
        raise ValueError("Use an HTTP(S) UI address without credentials")
    if not isinstance(qualification_notes, str) or len(qualification_notes.encode("utf-8")) > 32768:
        raise ValueError("Qualification notes must be at most 32 KiB of plain UTF-8 text")
    profiles = [
        " ".join(paragraph.split())
        for paragraph in re.split(r"\n\s*\n", qualification_notes)
        if paragraph.strip()
    ]
    if len(profiles) > 24 or any(len(paragraph) > 1200 for paragraph in profiles):
        raise ValueError("Use at most 24 qualification paragraphs, each at most 1200 characters")
    catalog = load_catalog(catalog_path)
    output.mkdir(parents=True, exist_ok=True)
    prefix = f"AI-Station-QA-{'DRAFT-' if draft else ''}{release[:12]}"
    destination = output / f"{prefix}.xlsx"
    csv_paths = {name: output / f"{prefix}-{name}.csv" for name in ("Testy", "Zgloszenia", "Retesty")}
    guide = output / f"{prefix}-Instrukcja.txt"
    manifest = output / f"{prefix}-manifest.json"
    fixtures = {output / path: data for path, data in fixture_bytes().items()}
    paths = [destination, *csv_paths.values(), guide, *fixtures, manifest]
    if any(path.exists() for path in paths):
        raise FileExistsError("Keep previous answers: choose a new output directory")
    catalog_hash = hashlib.sha256(
        json.dumps(catalog, sort_keys=True, ensure_ascii=False).encode()
    ).hexdigest()
    book = Workbook()
    book.remove(book.active)
    dark = PatternFill("solid", fgColor="18202B")
    editable = PatternFill("solid", fgColor="E8F3FF")
    wrap = Alignment(vertical="top", wrap_text=True)

    def sheet(name, headers, values, widths, freeze="A2"):
        page = book.create_sheet(name)
        page.append(headers)
        for row in values:
            page.append(row)
        for row in page:
            for cell in row:
                cell.alignment = wrap
                if isinstance(cell.value, str):
                    cell.data_type = "s"  # Catalog text must never become an Excel formula.
        for cell in page[1]:
            cell.fill = dark
            cell.font = Font(color="FFFFFF", bold=True)
        for index, width in enumerate(widths, 1):
            page.column_dimensions[get_column_letter(index)].width = width
        page.freeze_panes = freeze
        page.auto_filter.ref = page.dimensions
        page.sheet_view.zoomScale = 85
        page.row_dimensions[1].height = 34
        return page

    start = sheet(
        "Start",
        ["Informacja", "Instrukcja"],
        [
            [
                "Dokument",
                "WERSJA ROBOCZA — platforma nie jest jeszcze przekazana do testów operatora."
                if draft
                else "Arkusz testów operatora AI Station",
            ],
            ["Wersja aplikacji", release],
            ["Adres UI", base_url],
            ["Wygenerowano", datetime.now(UTC).isoformat(timespec="seconds")],
            ["Wersja katalogu", str(catalog["version"])],
            ["Identyfikator katalogu", catalog_hash],
            [
                "Kolejność",
                "Zacznij od uruchomienia i logowania, następnie jedna rozmowa, pliki i media, agenci, Memory oraz pozostałe obszary. Nie wybierasz trybu Chat / Work: orkiestrator dobiera wykonanie. P0 oznacza główną ścieżkę.",
            ],
            ["Gotowe dane i oczekiwane wartości", FIXTURE_INSTRUCTIONS],
            [
                "Zakres sprawdzonych profili",
                "Poniższe notatki opisują weryfikację wewnętrzną i jej zakres; nie są wynikami operatora. W testach zapisz rzeczywiście wybrany model i parametry."
                if profiles
                else "Opis sprawdzonych profili nie został dołączony. Uzupełnij go przed finalnym przekazaniem pakietu; nie zakładaj, że wszystkie modele i ustawienia zostały sprawdzone.",
            ],
            *[[f"Sprawdzony profil {index}", value] for index, value in enumerate(profiles, 1)],
            [
                "Co uzupełnić",
                "W Testy wypełniaj niebieskie pola. Wynik jest w kolumnie B i pozostaje widoczny podczas przewijania. Zapisz rzeczywisty rezultat, dowód oraz ID wykonania. Nie zmieniaj ID przypadków. Testy nie zostały za Ciebie oznaczone jako zaliczone.",
            ],
            [
                "Bezpieczne dane",
                "Używaj nowych plików, rozmów i projektów QA. Test Empty Trash wykonuj tylko, gdy cały kosz zawiera wyłącznie jednorazowe pliki QA. Nie usuwaj istniejących danych, aby zaliczyć przypadek.",
            ],
            [
                "Zasoby i konta",
                "GPU i zewnętrzne konta są opisane w każdym przypadku. Bez wymaganej konfiguracji wpisz Zablokowany i powód. Nie uruchamiaj płatnego API bez świadomej zgody na koszt.",
            ],
            [
                "Nie dotyczy",
                "Używaj tylko z uzasadnieniem. Brak obiecanej funkcji lub niedziałająca konfiguracja oznacza Zablokowany albo Błąd, a nie automatycznie Nie dotyczy.",
            ],
            [
                "Zgłoszenia",
                "Dla błędu lub pomysłu dodaj wiersz w Zgłoszenia. Wpisz ID testu, kroki i oczekiwane/rzeczywiste zachowanie. Nie wklejaj haseł, kluczy ani prywatnych dokumentów.",
            ],
            [
                "Retesty",
                "Po poprawce dodaj osobny wiersz Retesty z wersją poprawki i wynikiem. Nie nadpisuj oryginalnego błędu.",
            ],
            [
                "Zwrot wyników",
                "Zachowaj kopię, zapisz uzupełniony XLSX i przekaż go w rozmowie. Możesz też zwrócić CSV z zachowanymi ID. Wyniki służą do triage, poprawek i ponownych testów.",
            ],
            [
                "Znaczenie podsumowania",
                "Podsumowanie dotyczy wykonania tych testów; nie jest procentem ukończenia całego produktu. Formuły przeliczy Excel lub LibreOffice.",
            ],
        ],
        [24, 112],
    )
    start.auto_filter.ref = None
    for row in start.iter_rows(min_row=2):
        # Explicit height keeps wrapped instructions visible in spreadsheet viewers.
        width = 60 if str(row[0].value).startswith("Sprawdzony profil ") else 100
        start.row_dimensions[row[0].row].height = max(32, 16 * (len(wrap_text(str(row[1].value), width)) + 1))
    tests = sheet(
        "Testy",
        HEADERS,
        rows(catalog),
        [15, 20, 22, 38, 78, 68, 52, 52, 38, 44, 12, 60, 32, 10, 36, 36, 22, 24, 36],
        "E2",
    )
    for row in tests.iter_rows(min_row=2):
        lines = max(
            sum(
                max(
                    1,
                    len(
                        wrap_text(
                            paragraph,
                            max(8, int(tests.column_dimensions[cell.column_letter].width * 0.85)),
                        )
                    ),
                )
                for paragraph in str(cell.value or "").split("\n")
            )
            for cell in row
        )
        tests.row_dimensions[row[0].row].height = min(409, max(112, 16 * (lines + 1)))
        for index in [1, 6, 7, 8, 9, 14, 15, 16, 17]:
            row[index].fill = editable

    def choices(page, column, options, last):
        validation = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        validation.errorTitle = "Wybierz wynik z listy"
        validation.error = "Wybierz jedną z opisanych wartości. Szczegóły wpisz w uwagach."
        validation.showErrorMessage = True
        page.add_data_validation(validation)
        validation.add(f"{column}2:{column}{last}")

    last = len(catalog["cases"]) + 1
    choices(tests, "B", STATUSES, last)
    for status, color in [("OK", "DDF4DF"), ("Błąd", "FFDADA"), ("Zablokowany", "FFF0CE")]:
        tests.conditional_formatting.add(
            f"B2:B{last}",
            CellIsRule(operator="equal", formula=[f'"{status}"'], fill=PatternFill("solid", fgColor=color)),
        )
    capacity = max(100, len(catalog["cases"]) * 2)
    issues = sheet(
        "Zgłoszenia",
        ISSUE_HEADERS,
        [[f"ZGL-{i:03}", *([""] * 12)] for i in range(1, capacity + 1)],
        [18, 20, 24, 48, 72, 60, 60, 22, 20, 48, 48, 24, 36],
        "E2",
    )
    choices(issues, "B", ["Błąd", "Usprawnienie"], capacity + 1)
    choices(issues, "I", ["Krytyczny", "Wysoki", "Średni", "Niski"], capacity + 1)
    choices(
        issues, "L", ["Nowe", "Do wyjaśnienia", "Zaplanowane", "Naprawione", "Zweryfikowane"], capacity + 1
    )
    retests = sheet(
        "Retesty",
        RETEST_HEADERS,
        [[""] * 8 for _ in range(capacity)],
        [20, 20, 44, 22, 80, 24, 70, 50],
    )
    choices(retests, "F", STATUSES, capacity + 1)
    for page, first in [(issues, 2), (retests, 1)]:
        for row in page.iter_rows(min_row=2, min_col=first):
            for cell in row:
                cell.fill = editable
    summary = sheet(
        "Podsumowanie",
        ["Wskaźnik", "Liczba"],
        [["Liczba przypadków", len(catalog["cases"])], *[[status, ""] for status in STATUSES]],
        [36, 24],
    )
    for index, status in enumerate(STATUSES, 3):
        summary.cell(index, 2, f'=COUNTIF(Testy!B2:B{last},"{status}")')
    summary.append(["Obszar", "Liczba testów", "OK", "Błąd", "Zablokowany", "Nie testowano", "Nie dotyczy"])
    for area in dict.fromkeys(item["area"] for item in catalog["cases"]):
        index = summary.max_row + 1
        summary.append(
            [
                area,
                f"=COUNTIF(Testy!C2:C{last},A{index})",
                *[
                    f'=COUNTIFS(Testy!C2:C{last},A{index},Testy!B2:B{last},"{status}")'
                    for status in ("OK", "Błąd", "Zablokowany", "Nie testowano", "Nie dotyczy")
                ],
            ]
        )
    summary.auto_filter.ref = None
    for column in "CDEFG":
        summary.column_dimensions[column].width = 20
    for row in summary:
        for cell in row:
            cell.alignment = wrap
    sheet(
        "Słownik",
        ["Pojęcie", "Znaczenie"],
        [
            ["Nie testowano", "Nie wykonano jeszcze kroków; nie jest to wynik pozytywny."],
            ["OK", "Wykonano wszystkie kroki i uzyskano oczekiwany wynik."],
            ["Błąd", "Test wykonano, a rzeczywisty wynik różni się od oczekiwanego."],
            ["Zablokowany", "Nie można wykonać testu; zapisz brakujący warunek lub przeszkodę."],
            [
                "Nie dotyczy",
                "Świadomie poza uzgodnionym zakresem tej konfiguracji; zawsze wyjaśnij dlaczego.",
            ],
            ["Krytyczny / wysoki", "Utrata danych, naruszenie dostępu lub niedziałająca główna ścieżka."],
            ["Średni / niski", "Utrudniona funkcja z obejściem albo drobny problem prezentacji."],
            [
                "Usprawnienie",
                "Pomysł poprawiający użyteczność; oddziel go od niezgodności z obiecanym zachowaniem.",
            ],
            ["Run", "Pojedyncze wykonanie z własnym stanem, zdarzeniami, logami i wynikami."],
            ["Artefakt", "Zapisany wynik lub plik wejściowy z pochodzeniem i historią."],
        ],
        [30, 100],
    )
    with destination.open("xb") as handle:
        book.save(handle)
    for name, page in [("Testy", tests), ("Zgloszenia", issues), ("Retesty", retests)]:
        with csv_paths[name].open("x", encoding="utf-8-sig", newline="") as handle:
            writer = csv.writer(handle)
            for row in page.iter_rows(values_only=True):
                writer.writerow(
                    [
                        "'" + item
                        if isinstance(item, str) and item.lstrip().startswith(("=", "+", "-", "@"))
                        else item
                        for item in row
                    ]
                )
    with guide.open("x", encoding="utf-8") as handle:
        handle.write("AI STATION — QA OPERATORA\n\n")
        for key, value in start.iter_rows(min_row=2, values_only=True):
            handle.write(f"{key}\n{value}\n\n")
    for path, data in fixtures.items():
        path.parent.mkdir(parents=True, exist_ok=True)
        with path.open("xb") as handle:
            handle.write(data)
    with manifest.open("x", encoding="utf-8") as handle:
        json.dump(
            {
                "schema_version": 1,
                "release": release,
                "catalog_version": catalog["version"],
                "catalog_sha256": catalog_hash,
                "draft": draft,
                "base_url": base_url,
                "case_ids": [case["id"] for case in catalog["cases"]],
                "qualification_notes_sha256": hashlib.sha256("\n\n".join(profiles).encode()).hexdigest()
                if profiles
                else None,
                "qualification_paragraphs": len(profiles),
                "fixtures": [
                    {
                        "path": path.relative_to(output).as_posix(),
                        "size_bytes": len(data),
                        "sha256": hashlib.sha256(data).hexdigest(),
                    }
                    for path, data in fixtures.items()
                ],
            },
            handle,
            ensure_ascii=False,
            indent=2,
        )
    return paths


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--release", required=True)
    parser.add_argument("--base-url", default="http://127.0.0.1:8765")
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--catalog", type=Path, default=CATALOG)
    parser.add_argument("--draft", action="store_true")
    parser.add_argument("--qualification-notes-file", type=Path)
    args = parser.parse_args()
    notes = ""
    if args.qualification_notes_file:
        with args.qualification_notes_file.open(encoding="utf-8") as source:
            notes = source.read(32769)
    for path in export(
        args.output,
        release=args.release,
        base_url=args.base_url,
        draft=args.draft,
        catalog_path=args.catalog,
        qualification_notes=notes,
    ):
        print(path)


if __name__ == "__main__":
    main()
